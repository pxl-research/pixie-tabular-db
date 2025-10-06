from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from typing import Any, List

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class DocumentParseError(Exception):
    """Base class for document parsing issues."""


class UnsupportedDocumentType(DocumentParseError):
    """Raised when the uploaded document type is not supported."""


class EmptyDocumentError(DocumentParseError):
    """Raised when the uploaded document does not contain data rows."""


@dataclass(slots=True)
class ParsedDocument:
    source_name: str
    columns: List[str]
    rows: List[dict[str, Any]]


def parse_document(file_name: str, content: bytes) -> List[ParsedDocument]:
    lowered = file_name.lower()
    if lowered.endswith(".csv"):
        return [_parse_csv(file_name, content)]
    if lowered.endswith(".xlsx"):
        return _parse_xlsx(content)
    raise UnsupportedDocumentType("Only .csv and .xlsx files are supported at this stage")


def _parse_csv(file_name: str, content: bytes) -> ParsedDocument:
    try:
        text_stream = io.StringIO(content.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:  # pragma: no cover - depends on input file
        raise DocumentParseError("CSV file must be UTF-8 encoded") from exc

    reader = csv.DictReader(text_stream)
    if reader.fieldnames is None:
        raise EmptyDocumentError("CSV file is missing a header row")

    rows: List[dict[str, Any]] = []
    for raw_row in reader:
        row_dict = {field: raw_row.get(field) for field in reader.fieldnames}
        if all(value in (None, "") for value in row_dict.values()):
            continue
        rows.append(row_dict)

    if not rows:
        raise EmptyDocumentError("CSV file does not contain any data rows")

    base_name = file_name.rsplit("/", maxsplit=1)[-1]
    return ParsedDocument(
        source_name=base_name,
        columns=[field for field in reader.fieldnames],
        rows=rows,
    )


def _parse_xlsx(content: bytes) -> List[ParsedDocument]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except InvalidFileException as exc:
        raise DocumentParseError("XLSX file is invalid or corrupted") from exc

    parsed_sheets: List[ParsedDocument] = []

    try:
        for sheet in workbook.worksheets:
            if sheet is None:
                continue

            sheet_title = (sheet.title or "Sheet").strip() or "Sheet"
            rows_iter = sheet.iter_rows(values_only=True)

            try:
                header_row = next(rows_iter)
            except StopIteration:
                logging.info("Skipping empty sheet '%s'", sheet_title)
                continue

            columns: List[str] = []
            column_indexes: List[int] = []
            for idx, cell in enumerate(header_row):
                if cell is None:
                    continue
                column_name = str(cell).strip()
                if not column_name:
                    continue
                columns.append(column_name)
                column_indexes.append(idx)

            if not columns:
                logging.info("Skipping sheet '%s' without column names", sheet_title)
                continue

            rows: List[dict[str, Any]] = []
            for row in rows_iter:
                row_dict = {
                    columns[pos]: (
                        row[column_indexes[pos]] if column_indexes[pos] < len(row) else None
                    )
                    for pos in range(len(columns))
                }
                if all(value in (None, "") for value in row_dict.values()):
                    continue
                rows.append(row_dict)

            if not rows:
                logging.info("Skipping sheet '%s' without data rows", sheet_title)
                continue

            parsed_sheets.append(
                ParsedDocument(source_name=sheet_title, columns=columns, rows=rows)
            )
    finally:
        workbook.close()

    if not parsed_sheets:
        raise EmptyDocumentError("XLSX workbook does not contain any usable sheets")

    return parsed_sheets


__all__ = [
    "DocumentParseError",
    "EmptyDocumentError",
    "ParsedDocument",
    "UnsupportedDocumentType",
    "parse_document",
]
